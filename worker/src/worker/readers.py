"""Unified dataset readers: CSV / regular-tabular XLSX / Parquet.

All readers return a DataFrame with dtype=str (identifier-preserving contract)
after format-specific pre-checks. Pathological inputs are rejected with
explicit, actionable errors — never silently mis-aligned (F01).
"""

from __future__ import annotations

from pathlib import Path

import duckdb
import pandas as pd

SUPPORTED_EXTENSIONS = (".csv", ".xlsx", ".parquet")


class FormatError(ValueError):
    """Raised for unsupported or pathological input files."""


def _check_columns(columns: list[str], source: str) -> None:
    stripped = [str(c).strip() for c in columns]
    if len(stripped) == 0 or all(c == "" or c.lower() == "nan" for c in stripped):
        raise FormatError(f"{source}: empty header row")
    # pandas renders missing xlsx header cells as "Unnamed: N" placeholders.
    if any(c.startswith("Unnamed:") for c in stripped):
        raise FormatError(f"{source}: empty/missing header cells (pandas 'Unnamed:' placeholders)")
    if len(set(stripped)) != len(stripped):
        dupes = sorted({c for c in stripped if stripped.count(c) > 1})
        raise FormatError(f"{source}: duplicate column names: {', '.join(dupes)}")


def _finalize(df: pd.DataFrame) -> pd.DataFrame:
    df.columns = [str(c) for c in df.columns]
    return df.astype(str)


def read_csv(path: str) -> pd.DataFrame:
    df = pd.read_csv(path, dtype=str)
    _check_columns(list(df.columns), Path(path).name)
    return _finalize(df)


def read_xlsx(path: str, sheet_name: str | None = None) -> tuple[pd.DataFrame, list[str]]:
    from openpyxl import load_workbook

    # read_only mode omits merged_cells — normal mode is required for the check.
    book = load_workbook(path, read_only=False, data_only=True)
    try:
        sheet_names = list(book.sheetnames)
        target = sheet_name if sheet_name is not None else sheet_names[0]
        if target not in sheet_names:
            raise FormatError(
                f"{Path(path).name}: no sheet named '{target}'"
                f" (sheets: {', '.join(sheet_names)})"
            )
        sheet = book[target]
        if list(sheet.merged_cells.ranges):
            raise FormatError(
                f"{Path(path).name}!{target}: merged cells are not supported"
                " — flatten to a regular table first"
            )
        # Raw header check BEFORE pandas: pandas silently renames duplicate
        # columns (a -> a.1) and substitutes "Unnamed:" placeholders.
        header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        _check_columns(
            [str(h) if h is not None else "" for h in header],
            f"{Path(path).name}!{target}",
        )
    finally:
        book.close()

    df = pd.read_excel(path, dtype=str, sheet_name=target, engine="openpyxl")
    if len(df) == 0:
        raise FormatError(f"{Path(path).name}!{target}: sheet has no data rows")
    return _finalize(df), sheet_names


def read_parquet(path: str) -> pd.DataFrame:
    con = duckdb.connect(":memory:")
    try:
        con.execute("SET autoload_known_extensions = false")
        con.execute("SET autoinstall_known_extensions = false")
        described = con.execute("DESCRIBE SELECT * FROM read_parquet(?)", [path]).fetchall()
    except Exception as exc:  # noqa: BLE001 - surfaced as a format error with cause
        raise FormatError(f"{Path(path).name}: not readable as parquet ({exc})") from exc
    nested = [
        row[0]
        for row in described
        if str(row[1]).endswith("[]")
        or "STRUCT" in str(row[1]).upper()
        or "MAP" in str(row[1]).upper()
    ]
    if nested:
        raise FormatError(
            f"{Path(path).name}: nested columns not supported: {', '.join(nested)}"
        )
    df = con.execute("SELECT * FROM read_parquet(?)", [path]).df()
    _check_columns(list(df.columns), Path(path).name)
    return _finalize(df)


def read_any(path: str) -> pd.DataFrame:
    """Dispatch by extension; returns the dtype=str frame (CSV/XLSX/Parquet)."""
    lower = Path(path).suffix.lower()
    if lower == ".csv":
        return read_csv(path)
    if lower == ".xlsx":
        return read_xlsx(path)[0]
    if lower == ".parquet":
        return read_parquet(path)
    raise FormatError(
        f"{Path(path).name}: unsupported extension '{lower}'"
        f" (supported: {', '.join(SUPPORTED_EXTENSIONS)})"
    )


def sniff(path: str) -> dict:
    """Registration-time probe: format, row/column counts, sheet names (local-only info)."""
    lower = Path(path).suffix.lower()
    info: dict = {"format": lower.lstrip(".") if lower else "unknown"}
    if lower == ".xlsx":
        df, sheet_names = read_xlsx(path)
        info["sheet_names"] = sheet_names
        info["rows"] = int(len(df))
        info["columns"] = list(df.columns)
        if len(sheet_names) > 1:
            info["note"] = f"multiple sheets present; using first ('{sheet_names[0]}')"
    else:
        df = read_any(path)
        info["rows"] = int(len(df))
        info["columns"] = list(df.columns)
    return info
